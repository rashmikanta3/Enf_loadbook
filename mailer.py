import os
import io
import smtplib
from datetime import date
from typing import Optional, List, Dict
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.application import MIMEApplication
from dotenv import load_dotenv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

load_dotenv()

SMTP_SERVER = os.getenv("SMTP_SERVER", "smtp.gmail.com")
SMTP_PORT = int(os.getenv("SMTP_PORT", 465))
SENDER_EMAIL = os.getenv("SENDER_EMAIL")
SENDER_PASSWORD = os.getenv("SENDER_PASSWORD")

raw_recipients = os.getenv("MANAGEMENT_EMAIL", "")
MANAGEMENT_EMAILS = [email.strip() for email in raw_recipients.split(",") if email.strip()]


def create_excel_report(circle_summary: List[Dict], name_summary: List[Dict], from_date: Optional[date] = None, to_date: Optional[date] = None) -> bytes:
    wb = openpyxl.Workbook()

    header_fill_circle = PatternFill(start_color="0B3954", end_color="0B3954", fill_type="solid")
    header_fill_name = PatternFill(start_color="087E8B", end_color="087E8B", fill_type="solid")
    total_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    total_font = Font(name="Calibri", size=11, bold=True, color="000000")
    
    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )
    total_border = Border(
        left=Side(style='thin', color='A0AEC0'),
        right=Side(style='thin', color='A0AEC0'),
        top=Side(style='thin', color='718096'),
        bottom=Side(style='double', color='718096')
    )

    # ----------------------------------------------------
    # Sheet 1: Circle-Wise Summary
    # ----------------------------------------------------
    ws_circle = wb.active
    ws_circle.title = "Circle Summary"
    ws_circle.append(["Circle", "Total Cases", "Booked Load (kW)"])

    for col in range(1, 4):
        cell = ws_circle.cell(row=1, column=col)
        cell.fill = header_fill_circle
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    total_c_cases = 0
    total_c_kw = 0.0

    current_row = 2
    for row in circle_summary:
        cases = row.get("total_cases", 0)
        kw = round(float(row.get("total_load_kw", 0)), 2)
        total_c_cases += cases
        total_c_kw += kw

        ws_circle.append([row["circle"], cases, kw])
        for col_idx in range(1, 4):
            c = ws_circle.cell(row=current_row, column=col_idx)
            c.border = thin_border
            if col_idx > 1:
                c.alignment = Alignment(horizontal="right")
        current_row += 1

    # Total Row for Circle Summary
    ws_circle.append(["TOTAL", total_c_cases, round(total_c_kw, 2)])
    for col_idx in range(1, 4):
        c = ws_circle.cell(row=current_row, column=col_idx)
        c.fill = total_fill
        c.font = total_font
        c.border = total_border
        if col_idx > 1:
            c.alignment = Alignment(horizontal="right")

    # ----------------------------------------------------
    # Sheet 2: Name-Wise Summary (Sorted Circle-Wise)
    # ----------------------------------------------------
    ws_name = wb.create_sheet(title="Name Wise Summary")
    ws_name.append(["Circle", "Employee ID", "Employee Name", "Total Cases", "Booked Load (kW)"])

    for col in range(1, 6):
        cell = ws_name.cell(row=1, column=col)
        cell.fill = header_fill_name
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    total_n_cases = 0
    total_n_kw = 0.0

    current_row = 2
    for row in name_summary:
        cases = row.get("total_cases", 0)
        kw = round(float(row.get("total_load_kw", 0)), 2)
        total_n_cases += cases
        total_n_kw += kw

        ws_name.append([row["circle"], row["emp_id"], row["emp_name"], cases, kw])
        for col_idx in range(1, 6):
            c = ws_name.cell(row=current_row, column=col_idx)
            c.border = thin_border
            if col_idx >= 4:
                c.alignment = Alignment(horizontal="right")
        current_row += 1

    # Total Row for Name Summary
    ws_name.append(["TOTAL", "", "", total_n_cases, round(total_n_kw, 2)])
    for col_idx in range(1, 6):
        c = ws_name.cell(row=current_row, column=col_idx)
        c.fill = total_fill
        c.font = total_font
        c.border = total_border
        if col_idx >= 4:
            c.alignment = Alignment(horizontal="right")

    # Auto-adjust column widths
    for ws in [ws_circle, ws_name]:
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    excel_stream = io.BytesIO()
    wb.save(excel_stream)
    excel_stream.seek(0)
    return excel_stream.getvalue()


def send_summary_email(
    circle_summary: list, 
    name_summary: list, 
    from_date: Optional[date] = None, 
    to_date: Optional[date] = None
) -> bool:
    if not (SENDER_EMAIL and SENDER_PASSWORD and MANAGEMENT_EMAILS):
        print("SMTP credentials or management email recipients missing in .env.")
        return False

    # Ensure sorting: Circle ascending, and Name-wise by Circle then Employee Name
    sorted_circle_summary = sorted(circle_summary, key=lambda x: str(x.get("circle", "")).lower())
    sorted_name_summary = sorted(
        name_summary, 
        key=lambda x: (str(x.get("circle", "")).lower(), str(x.get("emp_name", "")).lower())
    )

    date_label = "All Time"
    if from_date and to_date:
        date_label = f"{from_date} to {to_date}"
    elif from_date:
        date_label = f"From {from_date}"
    elif to_date:
        date_label = f"Up to {to_date}"

    msg = MIMEMultipart("mixed")
    msg["Subject"] = f"Electricity Vigilance: Circle & Name Wise Summary ({date_label})"
    msg["From"] = SENDER_EMAIL
    msg["To"] = ", ".join(MANAGEMENT_EMAILS)

    # 1. Circle Table Rows & Total
    total_c_cases = sum(row.get("total_cases", 0) for row in sorted_circle_summary)
    total_c_kw = sum(float(row.get("total_load_kw", 0)) for row in sorted_circle_summary)

    circle_rows = "".join(
        f"<tr><td style='padding:6px 8px; border:1px solid #cbd5e0;'>{row['circle']}</td>"
        f"<td style='padding:6px 8px; border:1px solid #cbd5e0; text-align:center;'>{row['total_cases']}</td>"
        f"<td style='padding:6px 8px; border:1px solid #cbd5e0; text-align:right;'>{float(row['total_load_kw']):.2f} kW</td></tr>"
        for row in sorted_circle_summary
    )
    circle_rows += (
        f"<tr style='background-color:#edf2f7; font-weight:bold;'>"
        f"<td style='padding:7px 8px; border:1px solid #a0aec0;'>TOTAL</td>"
        f"<td style='padding:7px 8px; border:1px solid #a0aec0; text-align:center;'>{total_c_cases}</td>"
        f"<td style='padding:7px 8px; border:1px solid #a0aec0; text-align:right;'>{total_c_kw:.2f} kW</td>"
        f"</tr>"
    )

    # 2. Name Table Rows & Total (Circle as First Column)
    total_n_cases = sum(row.get("total_cases", 0) for row in sorted_name_summary)
    total_n_kw = sum(float(row.get("total_load_kw", 0)) for row in sorted_name_summary)

    name_rows = "".join(
        f"<tr>"
        f"<td style='padding:6px 8px; border:1px solid #cbd5e0;'><strong>{row['circle']}</strong></td>"
        f"<td style='padding:6px 8px; border:1px solid #cbd5e0; text-align:center;'>{row['emp_id']}</td>"
        f"<td style='padding:6px 8px; border:1px solid #cbd5e0;'>{row['emp_name']}</td>"
        f"<td style='padding:6px 8px; border:1px solid #cbd5e0; text-align:center;'>{row['total_cases']}</td>"
        f"<td style='padding:6px 8px; border:1px solid #cbd5e0; text-align:right;'>{float(row['total_load_kw']):.2f} kW</td>"
        f"</tr>"
        for row in sorted_name_summary
    )
    name_rows += (
        f"<tr style='background-color:#edf2f7; font-weight:bold;'>"
        f"<td colspan='3' style='padding:7px 8px; border:1px solid #a0aec0;'>TOTAL</td>"
        f"<td style='padding:7px 8px; border:1px solid #a0aec0; text-align:center;'>{total_n_cases}</td>"
        f"<td style='padding:7px 8px; border:1px solid #a0aec0; text-align:right;'>{total_n_kw:.2f} kW</td>"
        f"</tr>"
    )

    html_content = f"""
    <html>
    <body style="font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, Helvetica, Arial, sans-serif; color: #1a202c; line-height: 1.4;">
        <h2 style="color: #0b3954; margin-bottom: 4px;">Electricity Vigilance Inspection Summary</h2>
        <p style="margin-top: 0; color: #4a5568;"><strong>Inspection Period:</strong> {date_label}</p>
        <p style="font-size: 0.9rem; color: #718096;">The complete multi-sheet Excel report is attached below.</p>

        <h3 style="color: #0b3954; margin-bottom: 8px;">1. Circle-Wise Summary</h3>
        <table style="border-collapse: collapse; width: 100%; max-width: 550px; font-size: 13px;">
            <thead>
                <tr style="background:#0b3954; color:#ffffff;">
                    <th style="padding:8px; border:1px solid #0b3954; text-align:left;">Circle</th>
                    <th style="padding:8px; border:1px solid #0b3954; text-align:center; width: 90px;">Cases</th>
                    <th style="padding:8px; border:1px solid #0b3954; text-align:right; width: 140px;">Total Load</th>
                </tr>
            </thead>
            <tbody>{circle_rows}</tbody>
        </table>

        <br>
        <h3 style="color: #087e8b; margin-bottom: 8px;">2. Name-Wise Summary (Circle Sorted)</h3>
        <table style="border-collapse: collapse; width: 100%; max-width: 700px; font-size: 13px;">
            <thead>
                <tr style="background:#087e8b; color:#ffffff;">
                    <th style="padding:8px; border:1px solid #087e8b; text-align:left; width: 120px;">Circle</th>
                    <th style="padding:8px; border:1px solid #087e8b; text-align:center; width: 85px;">Emp ID</th>
                    <th style="padding:8px; border:1px solid #087e8b; text-align:left;">Employee Name</th>
                    <th style="padding:8px; border:1px solid #087e8b; text-align:center; width: 70px;">Cases</th>
                    <th style="padding:8px; border:1px solid #087e8b; text-align:right; width: 130px;">Total Load</th>
                </tr>
            </thead>
            <tbody>{name_rows}</tbody>
        </table>
    </body>
    </html>
    """
    msg.attach(MIMEText(html_content, "html"))

    excel_bytes = create_excel_report(sorted_circle_summary, sorted_name_summary, from_date, to_date)
    attachment_name = f"Vigilance_Summary_{date_label.replace(' ', '_')}.xlsx"
    part = MIMEApplication(excel_bytes, Name=attachment_name)
    part["Content-Disposition"] = f'attachment; filename="{attachment_name}"'
    msg.attach(part)

    try:
        with smtplib.SMTP_SSL(SMTP_SERVER, SMTP_PORT, timeout=20) as server:
            server.login(SENDER_EMAIL, SENDER_PASSWORD)
            server.sendmail(SENDER_EMAIL, MANAGEMENT_EMAILS, msg.as_string())
        print(f"Management report sent for period: {date_label}")
        return True
    except Exception as e:
        print(f"SMTP Error: {e}")
        return False